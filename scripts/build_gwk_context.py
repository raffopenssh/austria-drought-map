#!/usr/bin/env python3
"""Wasserschatz Oesterreichs (BMLUK 2021) groundwater-body (GWK) context.

Sources (committed):
  data/pop/wasserschatz_ergebnistabelle.xlsx
      Sheet '4 - GWK_Ressourcen': recharge & available resource per GWK
      (Mio m3/a, means 1998-2017), aquifer type, area.
      Sheet '5 - GWK_Wasserbedarf': water demand m3/a by sector
      (public supply / irrigation / livestock / industry / services),
      wells + springs, and Nutzungsintensitaet = demand/resource %.
  data/gwk/gwk.zip
      INSPIRE WFD GroundWaterBody NGP2015 GML (inspire.lfrz.gv.at 000801);
      138 features, localId GK1xxxxx; the 129 xlsx GWKs all match.

Outputs:
  web/data/gwk_context.json(.gz)
      gwks:  per GWK: name, aquifer type, area, resource, demand by
             sector, intensity %, population (allocated via KG shares),
             per-capita public-supply abstraction L/cap/d.
      kg2gwk: KG code -> GWK id (point-in-polygon of KG centroid;
             6 lakeside/border KGs use nearest body).
  web/data/gwk.geojson(.gz)
      simplified boundaries (localId, name, intensity) for map overlay.

Population allocation: each Gemeinde's latest population is split across
GWKs proportionally to how many of its KGs fall in each; Vienna handled
via the 90001 aggregate, merged Gemeinden via population.json alias map.
"""
import gzip
import json
import subprocess
import sys
import tempfile
import zipfile
from collections import defaultdict
from pathlib import Path

import openpyxl
from shapely.geometry import Point, shape
from shapely.strtree import STRtree

ROOT = Path(__file__).resolve().parent.parent
DATA = ROOT / "web/data"
XLSX = ROOT / "data/pop/wasserschatz_ergebnistabelle.xlsx"
GWK_ZIP = ROOT / "data/gwk/gwk.zip"

AQUIFER = {
    "PGWL": "porous aquifer",
    "vPGWL": "mostly porous aquifer",
    "vKAGWL": "mostly karst aquifer",
    "vKaGWL": "mostly karst aquifer",
    "vKLGWL": "mostly fissured aquifer",
    "vKlGWL": "mostly fissured aquifer",
}


def load_xlsx():
    wb = openpyxl.load_workbook(XLSX, read_only=True, data_only=True)
    out = {}
    for r in wb["4 - GWK_Ressourcen"].iter_rows(min_row=7, values_only=True):
        if not (r[0] and str(r[0]).startswith("GK")):
            continue
        out[str(r[0])] = {
            "name": str(r[1]).strip(),
            "group": str(r[2]) == "GWK-Gruppe",
            "aquifer": AQUIFER.get(str(r[3]).strip(), str(r[3]).strip()),
            "area_km2": round(float(r[4]), 1),
            "recharge_m3a": round(float(r[5]) * 1e6),   # Mio m3 -> m3
            "resource_note": (str(r[12]).strip() if r[12] else None),
        }
    for r in wb["5 - GWK_Wasserbedarf"].iter_rows(min_row=8, values_only=True):
        if not (r[0] and str(r[0]).startswith("GK")):
            continue
        g = out[str(r[0])]
        num = lambda v: int(v or 0)
        g["resource_m3a"] = num(r[3])
        # sector = wells + springs
        g["demand"] = {
            "supply": num(r[4]) + num(r[11]),
            "irrigation": num(r[5]) + num(r[12]),
            "livestock": num(r[6]) + num(r[13]),
            "industry": num(r[7]) + num(r[14]),
            "services": num(r[8]) + num(r[15]),
        }
        g["wells_m3a"] = num(r[9])
        g["springs_m3a"] = num(r[16])
        g["demand_m3a"] = g["wells_m3a"] + g["springs_m3a"]
        g["intensity_pct"] = round(float(r[10] or 0), 1)
        if r[17]:
            g["note"] = str(r[17]).strip()
    return out


def load_geoms(ids):
    with tempfile.TemporaryDirectory() as td:
        with zipfile.ZipFile(GWK_ZIP) as z:
            z.extractall(td)
        gml = next(Path(td).glob("*.gml"))
        gj = Path(td) / "gwk.geojson"
        subprocess.run(["ogr2ogr", "-f", "GeoJSON", "-t_srs", "EPSG:4326",
                        str(gj), str(gml)], check=True)
        feats = json.loads(gj.read_text())["features"]
    feats = [f for f in feats if f["properties"]["localId"] in ids]
    return {f["properties"]["localId"]: shape(f["geometry"]) for f in feats}


def main():
    gwks = load_xlsx()
    print(f"{len(gwks)} GWKs from xlsx")
    geoms = load_geoms(set(gwks))
    assert len(geoms) == len(gwks), (len(geoms), len(gwks))
    ids = list(geoms)
    tree = STRtree([geoms[i] for i in ids])

    kg_reg = json.loads((DATA / "kg_registry.json").read_text())
    kg2gwk, miss = {}, []
    for code, k in kg_reg.items():
        p = Point(k["lon"], k["lat"])
        hit = [ids[i] for i in tree.query(p) if geoms[ids[i]].contains(p)]
        if hit:
            kg2gwk[code] = hit[0]
        else:
            miss.append((code, p))
    for code, p in miss:  # lakes/border: nearest body
        kg2gwk[code] = min(ids, key=lambda i: geoms[i].distance(p))
    print(f"kg2gwk: {len(kg2gwk)} ({len(miss)} via nearest)")

    # population per GWK: split each Gemeinde by its KGs' GWK membership
    pop = json.loads((DATA / "population.json").read_text())
    gems, alias = pop["gemeinden"], pop.get("alias", {})
    latest = len(pop["years"]) - 1
    kgs_of_gem = defaultdict(list)
    for code, k in kg_reg.items():
        g = alias.get(k["g"], k["g"])
        kgs_of_gem[g].append(code)
    gwk_pop = defaultdict(float)
    allocated = 0
    for g, kgcodes in kgs_of_gem.items():
        rec = gems.get(g)
        if not rec:
            continue
        p = rec["t"][latest] / len(kgcodes)
        for code in kgcodes:
            gwk_pop[kg2gwk[code]] += p
        allocated += rec["t"][latest]
    print(f"population allocated: {allocated:,} of AT total")

    for gid, g in gwks.items():
        g["population"] = int(round(gwk_pop.get(gid, 0)))
        if g["population"] > 500:
            g["supply_lcd"] = round(
                g["demand"]["supply"] * 1000 / 365 / g["population"], 1)

    # national sanity: public-supply abstraction per capita
    tot_sup = sum(g["demand"]["supply"] for g in gwks.values())
    tot_pop = sum(g["population"] for g in gwks.values())
    print(f"AT public-supply GW abstraction: {tot_sup/1e6:.0f} Mio m3/a "
          f"= {tot_sup*1000/365/tot_pop:.0f} L/cap/d over {tot_pop:,} people")

    blob = json.dumps({
        "source": ("Wasserschatz Oesterreichs (BMLRT/Umweltbundesamt 2021), "
                   "Ergebnistabelle v02 2022-01; demand ~2017-2021, resource "
                   "means 1998-2017; GWK boundaries NGP2015 (INSPIRE)"),
        "note": ("population = Statistik Austria 2026, allocated to GWKs by "
                 "KG-share of each Gemeinde; supply_lcd = public-supply GW "
                 "abstraction per resident incl. losses/small business, vs "
                 "~130 L/cap/d household use (WAVE)"),
        "gwks": gwks, "kg2gwk": kg2gwk,
    }, separators=(",", ":"))
    (DATA / "gwk_context.json").write_text(blob)
    with gzip.open(DATA / "gwk_context.json.gz", "wt") as f:
        f.write(blob)
    print(f"gwk_context.json {len(blob)/1e6:.2f} MB")

    # simplified overlay
    with tempfile.TemporaryDirectory() as td:
        with zipfile.ZipFile(GWK_ZIP) as z:
            z.extractall(td)
        gml = next(Path(td).glob("*.gml"))
        raw = Path(td) / "raw.geojson"
        simp = Path(td) / "simp.geojson"
        subprocess.run(["ogr2ogr", "-f", "GeoJSON", "-t_srs", "EPSG:4326",
                        str(raw), str(gml)], check=True)
        subprocess.run(["ogr2ogr", "-f", "GeoJSON", str(simp), str(raw),
                        "-simplify", "0.002", "-select", "localId",
                        "-lco", "COORDINATE_PRECISION=4"], check=True)
        gj = json.loads(simp.read_text())
    feats = []
    for f in gj["features"]:
        gid = f["properties"]["localId"]
        g = gwks.get(gid)
        if not g:
            continue
        f["properties"] = {"id": gid, "n": g["name"],
                           "u": g["intensity_pct"]}
        feats.append(f)
    gj["features"] = feats
    blob = json.dumps(gj, separators=(",", ":"))
    (DATA / "gwk.geojson").write_text(blob)
    with gzip.open(DATA / "gwk.geojson.gz", "wt") as f:
        f.write(blob)
    print(f"gwk.geojson {len(blob)/1e6:.2f} MB ({len(feats)} bodies)")


if __name__ == "__main__":
    main()
